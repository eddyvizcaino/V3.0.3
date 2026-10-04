import io
import json
import os
import sqlite3
import tempfile
from pathlib import Path
from test_app import Base, PNG
from iglesia.db import init_db
from iglesia.portable_backup import export_backup, restore_empty

class TestV271(Base):
    def test_student_codes_are_automatic_and_unique_with_legacy_codes(self):
        cl = self.client('admin')
        self.post(cl, '/foligruc/students', {'FullName': 'Primera', 'MinistryId': '1', 'Code': 'manual'})
        self.assertEqual(self.q('SELECT Code FROM FStudents WHERE FullName=?', 'Primera')[0][0], 'FOL-000001')
        c = sqlite3.connect(self.db_path)
        c.execute("INSERT INTO FStudents(Code,FullName) VALUES('FOL-000009','Antigua')")
        c.commit(); c.close()
        self.post(cl, '/foligruc/students', {'FullName': 'Segunda', 'MinistryId': '1'})
        self.assertEqual(self.q('SELECT Code FROM FStudents WHERE FullName=?', 'Segunda')[0][0], 'FOL-000010')

    def test_foligruc_end_to_end_and_retakes(self):
        cl=self.client('admin')
        self.assertEqual(self.q('SELECT COUNT(*) FROM FSubjects')[0][0],15)
        self.assertEqual(self.q('SELECT Name FROM FSubjects WHERE Id=15')[0][0],'Misiología')
        for page in ['/foligruc/','/foligruc/students','/foligruc/subjects','/foligruc/periods','/foligruc/enrollments']:
            self.assertEqual(cl.get(page).status_code,200,page)
        self.post(cl,'/foligruc/students',{'FullName':'Ana Alumna','Code':'F-01','MinistryId':'1'})
        self.post(cl,'/foligruc/periods',{'Name':'2026 T1','StartDate':'2026-01-01'})
        self.assertEqual(self.q('SELECT EndDate FROM FPeriods')[0][0],'2026-03-31')
        self.post(cl,'/foligruc/enrollments',{'StudentId':1,'SubjectId':1,'PeriodId':1})
        self.post(cl,'/foligruc/students/1',{'RecordId':1,'Status':'Reprobada'})
        self.post(cl,'/foligruc/periods',{'Name':'2026 T2','StartDate':'2026-04-01'})
        self.post(cl,'/foligruc/enrollments',{'StudentId':1,'SubjectId':1,'PeriodId':2})
        self.post(cl,'/foligruc/students/1',{'RecordId':2,'Status':'Aprobada','ApprovedDate':'2026-06-30'})
        self.post(cl,'/foligruc/students/1',{'previous':'1','SubjectId':2})
        self.assertEqual(self.q('SELECT Status FROM FRecords ORDER BY Id'),[('Reprobada',),('Aprobada',),('Aprobada',)])
        self.assertGreater(self.q('SELECT COUNT(*) FROM FHistory')[0][0],3)
        self.assertIn('2 / 15',cl.get('/foligruc/students/1').text)
        self.post(cl,'/foligruc/students/1',{'previous':'1','SubjectId':2})
        self.assertEqual(self.q('SELECT COUNT(*) FROM FRecords')[0][0],3)
        self.assertTrue(cl.get('/foligruc/students/1/export/pdf').data.startswith(b'%PDF'))
        from openpyxl import load_workbook
        ws=load_workbook(io.BytesIO(cl.get('/foligruc/students/1/export/xlsx').data)).active
        self.assertEqual(ws['A3'].value,'Liderazgo Saludable')

    def test_foligruc_permissions_and_closed_period(self):
        admin=self.client('admin')
        self.post(admin,'/foligruc/students',{'FullName':'Protegido','Code':'F-1','MinistryId':'2'})
        limited=self.client('jovenes1')
        self.assertNotIn('Protegido',limited.get('/foligruc/students').text)
        for suffix in ('','/edit','/export/pdf','/export/xlsx'):
            self.assertEqual(limited.get('/foligruc/students/1'+suffix).status_code,403)
        self.assertEqual(self.post(limited,'/foligruc/periods',{'Name':'No','StartDate':'2026-01-01'}).status_code,403)
        self.post(admin,'/foligruc/periods',{'Name':'T1','StartDate':'2026-01-01'})
        self.post(admin,'/foligruc/enrollments',{'StudentId':1,'SubjectId':1,'PeriodId':1})
        self.post(admin,'/foligruc/periods',{'close':1})
        self.assertEqual(self.q('SELECT Closed FROM FPeriods')[0][0],0)
        self.post(admin,'/foligruc/students/1',{'RecordId':1,'Status':'Aprobada'})
        self.post(admin,'/foligruc/periods',{'close':1})
        self.assertEqual(self.post(admin,'/foligruc/students/1',{'RecordId':1,'Status':'Reprobada'}).status_code,400)

    def test_module_permissions_apply_immediately_and_user_creation(self):
        cl=self.client('admin');other=self.client('jovenes1')
        uid=self.q("SELECT Id FROM Users WHERE Username='jovenes1'")[0][0]
        self.post(cl,'/users',{'action':'permissions','uid':uid,'Permissions':['home','birthdays']})
        self.assertEqual(other.get('/members').status_code,403)
        self.assertEqual(other.get('/foligruc/').status_code,403)
        html=other.get('/').text
        self.assertNotIn('>MIEMBROS</a>',html)
        self.assertNotIn('>FOLIGRUC</a>',html)
        self.assertEqual(other.get('/birthdays').status_code,200)
        self.assertEqual(self.post(other,'/users',{'Username':'intruso'}).status_code,403)
        self.post(cl,'/users',{'Username':'secretaria','Password':'New-test-secret-1','Role':'Secretaria','Permissions':['home','foligruc']})
        self.assertEqual(self.q("SELECT Permissions FROM Users WHERE Username='secretaria'")[0][0],'home,foligruc')

    def test_private_photo_and_backup_roundtrip(self):
        cl=self.client('admin');self.member(cl,ProfilePhoto=(io.BytesIO(PNG),'photo.png'),content_type='multipart/form-data')
        # Base.member forwards content_type as a field; image tuple still creates multipart.
        response=cl.get('/members/1/photo');self.assertEqual(response.status_code,200)
        from PIL import Image
        self.assertEqual(Image.open(io.BytesIO(response.data)).size,(600,600))
        self.assertEqual(self.client('damas1').get('/members/1/photo').status_code,403)
        self.assertEqual(self.client().get('/members/1/photo').status_code,302)
        name=self.q('SELECT ProfilePhoto FROM Members')[0][0]
        self.assertEqual(cl.get('/static/profiles/'+name).status_code,404)
        blob=export_backup(self.db_path)
        new=str(Path(self.tmp)/'restored.db');init_db(new);restore_empty(new,blob)
        c=sqlite3.connect(new)
        self.assertEqual(c.execute('SELECT COUNT(*) FROM Members').fetchone()[0],1)
        self.assertEqual(c.execute('SELECT COUNT(*) FROM Photos').fetchone()[0],1);c.close()
        with self.assertRaises(ValueError):restore_empty(new,blob)
        corrupt=json.loads(blob);corrupt['tables']['Members'][0]['BogusColumn']='x'
        empty=str(Path(self.tmp)/'empty.db');init_db(empty)
        with self.assertRaises(ValueError):restore_empty(empty,json.dumps(corrupt))
        c=sqlite3.connect(empty);self.assertEqual(c.execute('SELECT COUNT(*) FROM FSubjects').fetchone()[0],15);c.close()

    def test_filters_calendar_and_no_logo(self):
        cl=self.client('admin');self.member(cl,FirstName='Filtrada',BirthDate='1990-01-01',BaptismDate='2020-01-01',MinistryId='1')
        self.member(cl,FirstName='Excluida',MinistryId='2')
        from openpyxl import load_workbook
        data=cl.get('/reports/members.xlsx?ministry=1&baptized=1&min_age=18').data
        ws=load_workbook(io.BytesIO(data)).active
        self.assertEqual(ws.max_row,2);self.assertIn('Filtrada',ws['A2'].value)
        self.assertIn('calendar-day',cl.get('/events?month=2026-09').text)
        self.assertNotIn('<img',self.client().get('/login').text)
        self.assertEqual(cl.get('/healthz').json['version'], self.app.config['VERSION'])

    def test_import_previous_academic_schema(self):
        from iglesia.legacy_import import convert_sqlite
        old=str(Path(self.tmp)/'old.db')
        c=sqlite3.connect(old)
        from iglesia.db import SCHEMA
        c.executescript(SCHEMA)
        c.executescript('''
        CREATE TABLE FSubjects(Id INTEGER PRIMARY KEY,Name TEXT,Position INTEGER,TeacherName TEXT);
        CREATE TABLE FPeriods(Id INTEGER PRIMARY KEY,Name TEXT,StartDate TEXT,EndDate TEXT,Status TEXT);
        CREATE TABLE FStudents(Id INTEGER PRIMARY KEY,Code TEXT,Name TEXT,MemberId INTEGER,EntryDate TEXT,Status TEXT);
        CREATE TABLE FEnrollments(Id INTEGER PRIMARY KEY,StudentId INTEGER,SubjectId INTEGER,PeriodId INTEGER,Status TEXT,FinalGrade REAL,CreatedAt TEXT,FinishedAt TEXT,TeacherName TEXT);
        CREATE TABLE FPriorCredits(Id INTEGER PRIMARY KEY,StudentId INTEGER,SubjectId INTEGER,ApprovedOn TEXT,FinalGrade REAL,CreatedBy INTEGER,CreatedAt TEXT,PeriodLabel TEXT);
        INSERT INTO Users(Id,Username,PasswordHash,Role) VALUES(1,'oldadmin','testhash','Admin');
        INSERT INTO Ministries(Id,Name) VALUES(1,'Ministerio');
        INSERT INTO Members(Id,FullName,MinistryId,CreatedAt) VALUES(1,'Persona heredada',1,'2020');
        INSERT INTO FSubjects VALUES(1,'Sanidad Interior',1,'Docente anterior'),(2,'Liderazgo Saludable',2,'');
        INSERT INTO FPeriods VALUES(1,'Anterior','2025-01-01','2025-03-31','Cerrado');
        INSERT INTO FStudents VALUES(1,'OLD-1','Estudiante heredado',1,'2024-01-01','Activo');
        INSERT INTO FEnrollments VALUES(1,1,1,1,'Reprobada',55,'2025-01-01','2025-03-31','Docente anterior');
        INSERT INTO FPriorCredits VALUES(1,1,2,NULL,NULL,1,'2025-01-01','2023');
        ''');c.commit();c.close()
        result=json.loads(convert_sqlite(old))['tables']
        self.assertEqual(result['FStudents'][0]['FullName'],'Estudiante heredado')
        self.assertEqual(result['FStudents'][0]['MinistryId'],1)
        self.assertEqual(result['FPeriods'][0]['Closed'],1)
        self.assertEqual([r['Status'] for r in result['FRecords']],['Reprobada','Aprobada'])
        self.assertIsNone(result['FRecords'][1]['Grade'])
        self.assertEqual(result['FRecords'][1]['Notes'],'2023')
        c=sqlite3.connect(old)
        self.assertNotIn('FullName',[r[1] for r in c.execute('PRAGMA table_info(FStudents)')]);c.close()
