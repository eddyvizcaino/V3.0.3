"""Pruebas automáticas. Ejecutar con:  python -m unittest discover -s tests -v"""
import io
import json
import os
import re
import shutil
import sqlite3
import sys
import tempfile
import time
import unittest

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from werkzeug.security import generate_password_hash  # noqa: E402

from iglesia import create_app  # noqa: E402
from iglesia import db as dbmod  # noqa: E402

from PIL import Image
_buffer = io.BytesIO()
Image.new('RGB', (20, 20), 'blue').save(_buffer, format='PNG')
PNG = _buffer.getvalue()
ADMIN_PW = 'Clave-Admin-9'


def token_from(html):
    m = re.search(r'name="csrf_token" value="([^"]+)"', html)
    return m.group(1) if m else ''


class Base(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.mkdtemp()
        self.uploads = os.path.join(self.tmp, 'profiles')
        self.app = create_app({
            'TESTING': True, 'SECRET_KEY': 'test',
            'DATABASE': os.path.join(self.tmp, 'iglesia.db'),
            'BACKUP_DIR': self.tmp, 'UPLOAD_DIR': self.uploads,
        })
        self.db_path = self.app.config['DATABASE']
        c = sqlite3.connect(self.db_path)
        c.execute('UPDATE Users SET PasswordHash=? WHERE Username=?', (generate_password_hash(ADMIN_PW), 'admin'))
        c.execute("INSERT INTO Ministries(Name) VALUES('Jóvenes'),('Damas')")
        for u, mid in (('jovenes1', 1), ('damas1', 2)):
            c.execute("INSERT INTO Users(Username,PasswordHash,FullName,Role,MinistryId,Active,CreatedAt) "
                      "VALUES(?,?,?,'Ministry',?,1,'x')", (u, generate_password_hash('Clave-Larga-1'), u, mid))
        c.commit()
        c.close()

    def tearDown(self):
        shutil.rmtree(self.tmp, ignore_errors=True)

    # ---- ayudas
    def q(self, sql, *p):
        c = sqlite3.connect(self.db_path)
        try:
            return c.execute(sql, p).fetchall()
        finally:
            c.close()

    def client(self, user=None, pw=None):
        cl = self.app.test_client()
        if user:
            self.login(cl, user, pw or ('Clave-Larga-1' if user != 'admin' else ADMIN_PW))
        return cl

    def login(self, cl, user, pw):
        tok = token_from(cl.get('/login').get_data(as_text=True))
        return cl.post('/login', data={'username': user, 'password': pw, 'csrf_token': tok})

    def post(self, cl, url, data=None, page='/', **kw):
        """POST con token CSRF tomado de una página cualquiera."""
        tok = token_from(cl.get(page).get_data(as_text=True))
        d = dict(data or {})
        d['csrf_token'] = tok
        return cl.post(url, data=d, **kw)

    def member(self, cl, **over):
        d = {'FirstName': 'Ana', 'LastName': 'Pérez', 'Status': 'Activo', 'GroupName': 'Damas',
             'MinistryId': '1'}
        d.update(over)
        return self.post(cl, '/members/new', d, page='/members/new', follow_redirects=True)


class TestAuth(Base):
    def test_fresh_db_and_default_password_forces_change(self):
        tmp = tempfile.mkdtemp()
        try:
            app = create_app({'TESTING': True, 'SECRET_KEY': 'k', 'DATABASE': os.path.join(tmp, 'n.db'),
                              'BACKUP_DIR': tmp, 'UPLOAD_DIR': os.path.join(tmp, 'p')})
            cl = app.test_client()
            r = self.login(cl, 'admin', 'admin123')
            self.assertEqual(r.headers['Location'], '/change-password')
            self.assertEqual(cl.get('/members').headers['Location'], '/change-password')
            tok = token_from(cl.get('/change-password').get_data(as_text=True))
            r = cl.post('/change-password', data={'current': 'admin123', 'new': 'admin123', 'csrf_token': tok})
            self.assertIn('inicial', r.get_data(as_text=True))          # rechaza seguir con la clave de fábrica
            r = cl.post('/change-password', data={'current': 'admin123', 'new': 'NuevaClave-77', 'csrf_token': tok})
            self.assertEqual(r.status_code, 302)
            self.assertEqual(cl.get('/members').status_code, 200)
        finally:
            shutil.rmtree(tmp, ignore_errors=True)

    def test_csrf_required(self):
        cl = self.client('admin')
        self.assertEqual(cl.post('/ministries/new', data={'Name': 'X'}).status_code, 400)
        self.assertEqual(self.q("select count(*) from Ministries where Name='X'")[0][0], 0)

    def test_login_bruteforce_lockout_and_reset(self):
        for i in range(5):
            self.login(self.app.test_client(), 'jovenes1', f'mala{i}')
        cl = self.app.test_client()
        r = self.login(cl, 'jovenes1', 'Clave-Larga-1')       # clave correcta pero en espera
        self.assertEqual(r.status_code, 200)
        self.assertIn('Demasiados intentos', r.get_data(as_text=True))
        self.assertEqual(self.login(self.app.test_client(), 'damas1', 'Clave-Larga-1').status_code, 302)  # otros no se afectan
        # tras el periodo de espera vuelve a entrar
        c = sqlite3.connect(self.db_path)
        c.execute("UPDATE LoginLog SET CreatedAt='2000-01-01T00:00:00'")
        c.commit()
        c.close()
        self.assertEqual(self.login(self.app.test_client(), 'jovenes1', 'Clave-Larga-1').status_code, 302)

    def test_locked_or_deactivated_user_loses_session_immediately(self):
        cl = self.client('jovenes1')
        self.assertEqual(cl.get('/members').status_code, 200)
        self.q('select 1')
        c = sqlite3.connect(self.db_path)
        c.execute("UPDATE Users SET Locked=1 WHERE Username='jovenes1'")
        c.commit()
        c.close()
        self.assertEqual(cl.get('/members').status_code, 302)

    def test_session_timeout(self):
        cl = self.client('admin')
        with cl.session_transaction() as s:
            s['last_activity'] = time.time() - 4000
        r = cl.get('/members')
        self.assertEqual(r.status_code, 302)
        self.assertIn('/login', r.headers['Location'])

    def test_logout_is_post_only(self):
        cl = self.client('admin')
        self.assertEqual(cl.get('/logout').status_code, 405)
        self.assertEqual(self.post(cl, '/logout').status_code, 302)
        self.assertEqual(cl.get('/members').status_code, 302)

    def test_cannot_lock_own_account(self):
        cl = self.client('admin')
        uid = self.q("select Id from Users where Username='admin'")[0][0]
        self.post(cl, '/security', {'uid': uid, 'locked': '1'}, page='/security')
        self.assertEqual(self.q('select Locked from Users where Id=?', uid)[0][0], 0)


class TestInjection(Base):
    def test_xss_and_template_injection_are_escaped(self):
        cl = self.client('admin')
        r = self.member(cl, FirstName='{{7*7}}', LastName='<script>alert(1)</script>', Cedula='"><b>x')
        html = r.get_data(as_text=True)
        self.assertNotIn('<script>alert(1)</script>', html)
        self.assertIn('&lt;script&gt;alert(1)&lt;/script&gt;', html)
        self.assertIn('{{7*7}}', html)              # se muestra literal, no se evalúa
        self.assertNotIn('>49<', html)
        form = cl.get('/members/1/edit').get_data(as_text=True)
        self.assertNotIn('"><b>x', form)

    def test_sql_metacharacters_in_search_params(self):
        cl = self.client('admin')
        self.assertEqual(cl.get("/members?ministry=1 OR 1=1").status_code, 200)
        self.assertEqual(cl.get("/audit?user=' OR '1'='1").status_code, 200)


class TestMembers(Base):
    def test_duplicate_cedula_friendly_error_and_no_db_lock(self):
        cl = self.client('admin')
        self.member(cl, Cedula='00100000001')
        self.assertEqual(self.q('select Cedula from Members')[0][0], '001-0000000-1')   # se normaliza
        r = self.member(cl, FirstName='Otra', Cedula='001-0000000-1')
        self.assertEqual(r.status_code, 200)
        self.assertIn('Ya existe un miembro registrado con esa cédula', r.get_data(as_text=True))
        self.assertIn('value="Otra"', r.get_data(as_text=True))             # conserva lo escrito
        self.assertEqual(self.q('select count(*) from Members')[0][0], 1)
        self.member(cl, FirstName='Tercera', Cedula='00200000002')            # la base sigue utilizable
        self.assertEqual(self.q('select count(*) from Members')[0][0], 2)

    def test_legacy_member_without_split_name_gets_suggestion(self):
        c = sqlite3.connect(self.db_path)
        c.execute("INSERT INTO Members(FullName,CreatedAt,MinistryId) VALUES('EDDY JOSE VIZCAINO PEREZ','x',1)")
        c.commit()
        c.close()
        html = self.client('admin').get('/members/1/edit').get_data(as_text=True)
        self.assertIn('value="EDDY JOSE"', html)
        self.assertIn('value="VIZCAINO PEREZ"', html)
        self.assertIn('registro es antiguo', html)

    def test_delete_requires_post_and_csrf(self):
        cl = self.client('admin')
        self.member(cl)
        self.assertEqual(cl.get('/members/1/delete').status_code, 405)
        self.assertEqual(cl.post('/members/1/delete').status_code, 400)
        self.assertEqual(self.q('select count(*) from Members')[0][0], 1)
        self.post(cl, '/members/1/delete', page='/members')
        self.assertEqual(self.q('select count(*) from Members WHERE DeletedAt IS NULL')[0][0], 0)

    def test_validation(self):
        cl = self.client('admin')
        r = self.member(cl, BirthDate='2999-01-01', Status='Hackeado', JoinDate='no-fecha')
        html = r.get_data(as_text=True)
        self.assertIn('no puede ser futura', html)
        self.assertIn('Estado: opción no válida', html)
        self.assertIn('Fecha de ingreso: fecha no válida', html)
        r = self.member(cl, MinistryId='999')
        self.assertIn('ministerio seleccionado no existe', r.get_data(as_text=True))
        self.assertEqual(self.q('select count(*) from Members')[0][0], 0)

    def test_edit_history_lists_changed_fields(self):
        cl = self.client('admin')
        self.member(cl, Phone='111')
        self.post(cl, '/members/1/edit', {'FirstName': 'Ana', 'LastName': 'Pérez', 'Phone': '222',
                                          'Status': 'Activo', 'GroupName': 'Damas', 'MinistryId': '1',
                                          'ChurchRole': 'Miembro', 'FollowUpStatus': 'Activo'},
                  page='/members/1/edit')
        detail = self.q("select Details from MemberHistory where Action='Editar'")[0][0]
        self.assertEqual(detail, 'Cambios: Teléfono')

    def test_photo_validation_and_cleanup(self):
        cl = self.client('admin')
        base = {'FirstName': 'Ana', 'LastName': 'P', 'Status': 'Activo', 'GroupName': 'Damas', 'MinistryId': '1'}
        fake = dict(base, ProfilePhoto=(io.BytesIO(b'<?php evil ?>'), 'x.jpg'))
        r = self.post(cl, '/members/new', fake, page='/members/new', content_type='multipart/form-data',
                      follow_redirects=True)
        self.assertIn('imagen JPG, PNG o WEBP válida', r.get_data(as_text=True))
        self.assertEqual(os.listdir(self.uploads), [])
        big = dict(base, ProfilePhoto=(io.BytesIO(PNG + b'0' * (5 * 1024 * 1024)), 'big.png'))
        r = self.post(cl, '/members/new', big, page='/members/new', content_type='multipart/form-data',
                      follow_redirects=True)
        self.assertIn('supera los 5 MB', r.get_data(as_text=True))
        ok = dict(base, ProfilePhoto=(io.BytesIO(PNG), 'a.png'))
        self.post(cl, '/members/new', ok, page='/members/new', content_type='multipart/form-data')
        first = self.q('select ProfilePhoto from Members')[0][0]
        self.assertTrue(first.endswith('.jpg') and bool(self.q('SELECT Id FROM Photos WHERE Name=?',first)))
        # reemplazar borra la anterior
        new = dict(base, ProfilePhoto=(io.BytesIO(PNG), 'b.png'), ChurchRole='Miembro', FollowUpStatus='Activo')
        self.post(cl, '/members/1/edit', new, page='/members/1/edit', content_type='multipart/form-data')
        second = self.q('select ProfilePhoto from Members')[0][0]
        self.assertNotEqual(first, second)
        self.assertEqual(self.q('SELECT Name FROM Photos'), [(second,)])
        # enviar a papelera conserva la foto para restaurar
        self.post(cl, '/members/1/delete', page='/members')
        self.assertEqual(self.q('SELECT Name FROM Photos'), [(second,)])


class TestPermissions(Base):
    def setUp(self):
        super().setUp()
        adm = self.client('admin')
        self.member(adm, FirstName='DeJovenes', MinistryId='1', Cedula='00100000001')
        self.member(adm, FirstName='DeDamas', MinistryId='2', Cedula='00200000002')

    def test_ministry_user_sees_only_own_members(self):
        html = self.client('jovenes1').get('/members').get_data(as_text=True)
        self.assertIn('DeJovenes', html)
        self.assertNotIn('DeDamas', html)

    def test_ministry_user_cannot_touch_other_ministry_members(self):
        cl = self.client('jovenes1')
        for url in ('/members/2', '/members/2/edit'):
            r = cl.get(url)
            self.assertEqual(r.status_code, 302, url)
        self.post(cl, '/members/2/delete', page='/members')
        self.assertEqual(self.q('select count(*) from Members')[0][0], 2)
        r = self.post(cl, '/members/2/edit', {'FirstName': 'Hack', 'LastName': 'X', 'Status': 'Activo',
                                              'GroupName': 'Damas'}, page='/members')
        self.assertEqual(self.q('select FirstName from Members where Id=2')[0][0], 'DeDamas')

    def test_ministry_id_tampering_is_ignored(self):
        cl = self.client('jovenes1')
        self.member(cl, FirstName='Nuevo', MinistryId='2', Cedula='00300000003')
        self.assertEqual(self.q("select MinistryId from Members where FirstName='Nuevo'")[0][0], 1)

    def test_admin_only_pages_and_reports(self):
        cl = self.client('jovenes1')
        for url in ('/users', '/security', '/audit', '/backup', '/reports/complete.xlsx', '/ministries/new'):
            r = cl.get(url)
            self.assertIn(r.status_code, (302,403), url)
        self.assertEqual(self.post(cl, '/ministries/1/delete', page='/ministries').status_code, 302)
        self.assertEqual(self.q('select count(*) from Ministries')[0][0], 2)

    def test_ministry_without_assignment_sees_nothing_and_cannot_create(self):
        c = sqlite3.connect(self.db_path)
        c.execute("UPDATE Users SET MinistryId=NULL WHERE Username='jovenes1'")
        c.commit()
        c.close()
        cl = self.client('jovenes1')
        self.assertNotIn('DeJovenes', cl.get('/members').get_data(as_text=True))
        self.assertEqual(cl.get('/members/1').status_code, 302)
        r = self.member(cl, FirstName='Huerfano')
        self.assertIn('no tiene un ministerio asignado', r.get_data(as_text=True))


class TestAdmin(Base):
    def test_user_creation_rules(self):
        cl = self.client('admin')
        def create(**o):
            d = {'Username': 'nuevo', 'Password': 'Clave-Buena-5', 'Role': 'Ministry', 'MinistryId': '1'}
            d.update(o)
            return self.post(cl, '/users', d, page='/users', follow_redirects=True).get_data(as_text=True)
        self.assertIn('al menos 8', create(Password='corta'))
        self.assertIn('debe tener un ministerio', create(MinistryId=''))
        self.assertIn('Tipo de usuario no válido', create(Role='Root'))
        self.assertIn('Usuario creado', create())
        self.assertIn('ya existe', create())
        self.assertEqual(self.q("select MinistryId from Users where Username='nuevo'")[0][0], 1)

    def test_ministry_delete_blocked_when_in_use(self):
        cl = self.client('admin')
        self.member(cl, MinistryId='1')
        r = self.post(cl, '/ministries/1/delete', page='/ministries', follow_redirects=True)
        self.assertIn('No se puede borrar', r.get_data(as_text=True))
        c = sqlite3.connect(self.db_path)
        c.execute("INSERT INTO Ministries(Name) VALUES('Vacío')")
        c.commit()
        c.close()
        self.post(cl, '/ministries/3/delete', page='/ministries')
        self.assertEqual(self.q("select count(*) from Ministries where Name='Vacío'")[0][0], 0)

    def test_backup_and_restore_roundtrip(self):
        cl = self.client('admin')
        self.member(cl, FirstName='Original')
        r = self.post(cl, '/backup/download', page='/backup')
        self.assertEqual(r.status_code, 200)
        name = [n for n in os.listdir(self.tmp) if n.startswith('iglesia_backup_')][0]
        self.post(cl, '/members/1/delete', page='/members')
        self.assertEqual(self.q('select count(*) from Members WHERE DeletedAt IS NULL')[0][0], 0)
        r = self.post(cl, '/backup/restore', {'saved_backup': name}, page='/backup')
        self.assertEqual(r.status_code, 302)
        self.assertEqual(self.q('select FirstName from Members')[0][0], 'Original')
        self.assertTrue([n for n in os.listdir(self.tmp) if n.startswith('iglesia_ANTES_RESTAURAR_')])
        self.assertEqual(cl.get('/members').status_code, 302)                  # sesión cerrada tras restaurar
        self.assertTrue(self.q("select 1 from AuditLog where Action='Restaurar'"))

    def test_restore_works_when_backup_lacks_the_current_admin(self):
        """Antes fallaba al registrar la auditoría si el usuario actual no existía en la copia."""
        cl0 = self.client('admin')
        self.post(cl0, '/backup/download', page='/backup')                     # copia SIN el admin2
        name = [n for n in os.listdir(self.tmp) if n.startswith('iglesia_backup_')][0]
        c = sqlite3.connect(self.db_path)
        c.execute("INSERT INTO Users(Username,PasswordHash,FullName,Role,Active,CreatedAt) "
                  "VALUES('admin2',?, 'A2','Admin',1,'x')", (generate_password_hash('Clave-Larga-1'),))
        c.commit()
        c.close()
        cl = self.client('admin2', 'Clave-Larga-1')
        r = self.post(cl, '/backup/restore', {'saved_backup': name}, page='/backup', follow_redirects=True)
        self.assertIn('Copia restaurada correctamente', r.get_data(as_text=True))
        self.assertEqual(self.q("select count(*) from Users where Username='admin2'")[0][0], 0)
        self.assertIn('admin2', self.q("select Details from AuditLog where Action='Restaurar'")[0][0])

    def test_restore_rejects_bad_files_and_path_traversal(self):
        cl = self.client('admin')
        for saved in ('../iglesia.db', 'app.py', '..\\x.db', 'iglesia.db'):
            r = self.post(cl, '/backup/restore', {'saved_backup': saved}, page='/backup', follow_redirects=True)
            self.assertIn('no válida', r.get_data(as_text=True), saved)
        junk = {'backup_file': (io.BytesIO(b'esto no es sqlite'), 'x.db')}
        r = self.post(cl, '/backup/restore', junk, page='/backup', content_type='multipart/form-data',
                      follow_redirects=True)
        self.assertIn('no es una base de datos SQLite', r.get_data(as_text=True))
        other = sqlite3.connect(os.path.join(self.tmp, 'otra.db'))
        other.execute('create table Foo(x)')
        other.commit()
        other.close()
        with open(os.path.join(self.tmp, 'otra.db'), 'rb') as fh:
            r = self.post(cl, '/backup/restore', {'backup_file': (fh, 'otra.db')}, page='/backup',
                          content_type='multipart/form-data', follow_redirects=True)
        self.assertIn('no parece ser una base de datos válida', r.get_data(as_text=True))
        self.assertEqual(self.q('select count(*) from Users')[0][0], 3)         # base intacta

    def test_autobackup_keeps_daily_and_skips_corrupt_db(self):
        dbmod.run_autobackup_once(self.app)
        auto = os.path.join(self.tmp, dbmod.AUTOBACKUP_NAME)
        self.assertTrue(os.path.exists(auto))
        self.assertTrue([n for n in os.listdir(self.tmp) if n.startswith('iglesia_backup_diario_')])
        good = os.path.getsize(auto)
        with open(self.db_path, 'r+b') as fh:                                   # corrompe la base
            fh.seek(100)
            fh.write(b'\xff' * 4000)
        dbmod.run_autobackup_once(self.app)
        self.assertEqual(os.path.getsize(auto), good)                           # la copia buena se conserva


class TestEvents(Base):
    def new_event(self, cl, **over):
        d = {'Name': 'Culto especial', 'EventDate': '2999-05-01', 'EventTime': '19:30', 'Venue': 'Templo',
             'Status': 'Próximo', 'MinistryId': ''}
        d.update(over)
        return self.post(cl, '/events/new', d, page='/events/new', follow_redirects=True)

    def setUp(self):
        super().setUp()
        adm = self.client('admin')
        self.new_event(adm, Name='GeneralEv', MinistryId='')
        self.new_event(adm, Name='DeJovenesEv', MinistryId='1')
        self.new_event(adm, Name='DeDamasEv', MinistryId='2')

    def test_visibility_general_plus_own_ministry(self):
        html = self.client('jovenes1').get('/events').get_data(as_text=True)
        self.assertIn('GeneralEv', html)
        self.assertIn('DeJovenesEv', html)
        self.assertNotIn('DeDamasEv', html)
        home = self.client('jovenes1').get('/').get_data(as_text=True)
        self.assertIn('GeneralEv', home)
        self.assertNotIn('DeDamasEv', home)

    def test_ministry_user_permissions(self):
        cl = self.client('jovenes1')
        ids = {n: i for i, n in self.q('select Id,Name from Events')}
        # solo puede editar/eliminar los de su ministerio
        for name in ('GeneralEv', 'DeDamasEv'):
            self.assertEqual(cl.get(f'/events/{ids[name]}/edit').status_code, 302, name)
            self.post(cl, f'/events/{ids[name]}/delete', page='/events')
            self.assertEqual(self.q('select count(*) from Events where Name=?', name)[0][0], 1, name)
        self.assertEqual(cl.get(f'/events/{ids["DeJovenesEv"]}/edit').status_code, 200)
        self.post(cl, f'/events/{ids["DeJovenesEv"]}/edit',
                  {'Name': 'Renombrado', 'EventDate': '2999-06-01', 'Status': 'Próximo', 'MinistryId': '2'},
                  page='/events')
        name, mid = self.q('select Name,MinistryId from Events where Id=?', ids['DeJovenesEv'])[0]
        self.assertEqual((name, mid), ('Renombrado', 1))                      # no puede cambiarlo de ministerio
        self.post(cl, f'/events/{ids["DeJovenesEv"]}/delete', page='/events')
        self.assertEqual(self.q('select count(*) from Events')[0][0], 2)

    def test_ministry_user_creates_only_in_own_ministry(self):
        self.new_event(self.client('jovenes1'), Name='MiEvento', MinistryId='2')
        self.assertEqual(self.q("select MinistryId from Events where Name='MiEvento'")[0][0], 1)

    def test_validation(self):
        cl = self.client('admin')
        html = self.new_event(cl, Name='', EventDate='', EventTime='25:99', Status='Raro',
                              MinistryId='99').get_data(as_text=True)
        for msg in ('título del evento es obligatorio', 'fecha del evento es obligatoria', 'formato HH:MM',
                    'Estado: opción no válida', 'ministerio seleccionado no existe'):
            self.assertIn(msg, html)
        self.assertEqual(self.q('select count(*) from Events')[0][0], 3)

    def test_delete_needs_post_and_csrf_and_is_audited(self):
        cl = self.client('admin')
        self.assertEqual(cl.get('/events/1/delete').status_code, 405)
        self.assertEqual(cl.post('/events/1/delete').status_code, 400)
        self.post(cl, '/events/1/delete', page='/events')
        self.assertEqual(self.q('select count(*) from Events where Id=1')[0][0], 0)
        self.assertTrue(self.q("select 1 from AuditLog where Entity='Evento' and Action='Eliminar'"))

    def test_edit_records_changes_and_escapes_html(self):
        cl = self.client('admin')
        self.post(cl, '/events/1/edit', {'Name': '<b>x</b>', 'EventDate': '2999-05-01', 'EventTime': '19:30',
                                         'Venue': 'Templo', 'Status': 'Cancelado', 'MinistryId': ''},
                  page='/events')
        self.assertIn('Cambios: Título, Estado', self.q("select Details from AuditLog where Action='Editar' and Entity='Evento'")[0][0])
        self.assertIn('&lt;b&gt;x&lt;/b&gt;', cl.get('/events').get_data(as_text=True))

    def test_legacy_status_is_preserved_when_editing(self):
        c = sqlite3.connect(self.db_path)
        c.execute("INSERT INTO Events(Name,EventDate,Status,CreatedAt) VALUES('Viejo','2001-01-01','En curso','x')")
        c.commit()
        c.close()
        cl = self.client('admin')
        eid = self.q("select Id from Events where Name='Viejo'")[0][0]
        self.assertIn('<option selected>En curso</option>', cl.get(f'/events/{eid}/edit').get_data(as_text=True))

    def test_upcoming_listed_before_past(self):
        c = sqlite3.connect(self.db_path)
        c.execute("INSERT INTO Events(Name,EventDate,CreatedAt) VALUES('PasadoEv','2001-01-01','x')")
        c.commit()
        c.close()
        html = self.client('admin').get('/events').get_data(as_text=True)
        self.assertLess(html.index('GeneralEv'), html.index('PasadoEv'))


class TestReports(Base):
    def test_exports(self):
        import openpyxl
        cl = self.client('admin')
        self.member(cl, FirstName='=CMD()', LastName='Peligro', Cedula='00100000001')
        r = cl.get('/reports/members.xlsx')
        self.assertEqual(r.status_code, 200)
        ws = openpyxl.load_workbook(io.BytesIO(r.data)).active
        self.assertEqual(ws['A2'].value, "'=CMD() Peligro")                      # no se ejecuta como fórmula
        self.assertTrue(cl.get('/reports/members.pdf').data.startswith(b'%PDF'))
        r = cl.get('/reports/complete.xlsx')
        self.assertEqual(openpyxl.load_workbook(io.BytesIO(r.data)).sheetnames, ['Miembros', 'Ministerios', 'Actividad'])

    def test_ministry_cannot_export_admin_reports(self):
        for endpoint in ('/reports','/reports/members.xlsx','/reports/members.pdf','/reports/complete.xlsx'):
            self.assertEqual(self.client('jovenes1').get(endpoint).status_code,403)


class TestPages(Base):
    def test_home_only_lists_future_events(self):
        c = sqlite3.connect(self.db_path)
        c.execute("INSERT INTO Events(Name,EventDate,CreatedAt) VALUES('EventoViejo','2001-01-01','x')")
        c.execute("INSERT INTO Events(Name,EventDate,CreatedAt) VALUES('EventoFuturo','2999-01-01','x')")
        c.commit()
        c.close()
        html = self.client('admin').get('/').get_data(as_text=True)
        self.assertIn('EventoFuturo', html)
        self.assertNotIn('EventoViejo', html)

    def test_every_page_renders_for_admin_and_ministry(self):
        adm = self.client('admin')
        self.member(adm, BirthDate='1990-05-05', Cedula='00100000001')
        urls = ['/', '/members', '/members/new', '/members/1', '/members/1/edit', '/ministries',
                '/ministries/new', '/ministries/1/edit', '/events', '/events/new', '/alerts', '/birthdays', '/reports',
                '/change-password', '/users', '/audit', '/security', '/backup']
        for u in urls:
            self.assertEqual(adm.get(u).status_code, 200, u)
        mini = self.client('jovenes1')
        for u in ['/', '/members', '/members/new', '/ministries', '/events', '/birthdays', '/foligruc/']:
            self.assertEqual(mini.get(u).status_code, 200, u)
        self.assertEqual(adm.get('/no-existe').status_code, 404)

    def test_alert_creation_and_validation(self):
        cl = self.client('admin')
        r = self.post(cl, '/alerts', {'Title': '', 'Priority': 'Normal'}, page='/alerts', follow_redirects=True)
        self.assertIn('título es obligatorio', r.get_data(as_text=True))
        r = self.post(cl, '/alerts', {'Title': 'Reunión <b>', 'Priority': 'Alta', 'DueDate': '2026-10-01'},
                      page='/alerts', follow_redirects=True)
        html = r.get_data(as_text=True)
        self.assertIn('Reunión &lt;b&gt;', html)

    def test_security_headers(self):
        r = self.client('admin').get('/')
        self.assertEqual(r.headers['X-Frame-Options'], 'DENY')
        self.assertEqual(r.headers['Cache-Control'], 'no-store')


class TestMigration(unittest.TestCase):
    LEGACY = '/mnt/user-data/uploads/iglesia.db'

    def test_replace_backup_is_atomic_and_keeps_user_accounts(self):
        from iglesia.portable_backup import export_backup, restore_empty
        tmp = tempfile.mkdtemp()
        try:
            path = os.path.join(tmp, 'iglesia.db')
            dbmod.init_db(path)
            c = sqlite3.connect(path)
            c.execute("INSERT INTO Members(FullName,CreatedAt) VALUES('Ana original','2026-01-01')")
            c.commit(); c.close()
            good = export_backup(path)
            c = sqlite3.connect(path)
            c.execute("UPDATE Members SET FullName='Cambio posterior'")
            c.commit(); c.close()
            broken = json.loads(good)
            broken['tables']['Members'][0]['CampoInexistente'] = 'x'
            with self.assertRaises(ValueError):
                restore_empty(path, json.dumps(broken), replace_existing=True)
            c = sqlite3.connect(path)
            self.assertEqual(c.execute('SELECT FullName FROM Members').fetchone()[0], 'Cambio posterior')
            c.close()
            restore_empty(path, good, replace_existing=True)
            c = sqlite3.connect(path)
            self.assertEqual(c.execute('SELECT FullName FROM Members').fetchone()[0], 'Ana original')
            self.assertEqual(c.execute('SELECT COUNT(*) FROM Users').fetchone()[0], 1)
            c.close()
        finally:
            shutil.rmtree(tmp, ignore_errors=True)

    def test_old_users_can_log_in_after_migration(self):
        tmp = tempfile.mkdtemp()
        try:
            path = os.path.join(tmp, 'iglesia.db')
            old = sqlite3.connect(path)
            old.execute('CREATE TABLE Users (Id INTEGER PRIMARY KEY, Username TEXT UNIQUE, PasswordHash TEXT, Role TEXT)')
            old.execute('INSERT INTO Users VALUES (1,?,?,?)',
                        ('oldadmin', generate_password_hash('ClaveSegura123!'), 'Admin'))
            old.commit(); old.close()
            dbmod.init_db(path)
            app = create_app({'TESTING': True, 'DATABASE': path,
                              'SECRET_KEY': 'test-key', 'UPLOAD_DIR': os.path.join(tmp, 'profiles'),
                              'BACKUP_DIR': tmp})
            with app.test_client() as client:
                token = token_from(client.get('/login').get_data(as_text=True))
                response = client.post('/login', data={'username': 'oldadmin',
                    'password': 'ClaveSegura123!', 'csrf_token': token}, follow_redirects=True)
                self.assertEqual(response.status_code, 200)
            old = sqlite3.connect(path)
            self.assertEqual(old.execute('SELECT Locked,FailedAttempts,Active FROM Users WHERE Id=1').fetchone(), (0, 0, 1))
            old.close()
        finally:
            shutil.rmtree(tmp, ignore_errors=True)

    def test_legacy_database_is_upgraded_without_losing_data(self):
        tmp = tempfile.mkdtemp()
        try:
            path = os.path.join(tmp, 'iglesia.db')
            legacy=sqlite3.connect(path)
            legacy.executescript(dbmod.SCHEMA)
            legacy.execute("INSERT INTO Members(FullName,CreatedAt) VALUES('Miembro heredado','2020-01-01')")
            legacy.execute("INSERT INTO Users(Username,PasswordHash,Role) VALUES('anterior','hash-sintetico','Admin')")
            legacy.commit();legacy.close()
            before = sqlite3.connect(path)
            n_members = before.execute('select count(*) from Members').fetchone()[0]
            n_users = before.execute('select count(*) from Users').fetchone()[0]
            before.close()
            dbmod.init_db(path)
            dbmod.init_db(path)  # idempotente
            c = sqlite3.connect(path)
            cols = {r[1] for r in c.execute('pragma table_info(Members)')}
            for col in ('FollowUpStatus', 'ChurchRole', 'BaptismDate', 'ProfilePhoto'):
                self.assertIn(col, cols)
            self.assertEqual(c.execute('select count(*) from Members').fetchone()[0], n_members)
            self.assertEqual(c.execute('select count(*) from Users').fetchone()[0], n_users)
            self.assertEqual(c.execute('pragma integrity_check').fetchone()[0], 'ok')
            c.close()
        finally:
            shutil.rmtree(tmp, ignore_errors=True)


if __name__ == '__main__':
    unittest.main()
